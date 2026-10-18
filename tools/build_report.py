"""Build the APA student paper from documented code and the saved real benchmark.

Requires python-docx. Does not run or replace the benchmark. PDF export is a
separate step (export_report.py). Word pagination can be checked independently.
"""

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
TITLE = "Hash Tables and Binary Max-Heaps: Design, Verification, and Performance"
REFERENCES = [
    ("Black, P. E. (2019, February 12). Perfect hashing. In Dictionary of Algorithms and Data Structures. National Institute of Standards and Technology. ",
     "https://xlinux.nist.gov/dads/HTML/perfecthash.html"),
    ("Morin, P. (n.d.-a). BinaryHeap: An implicit binary tree. In Open data structures (in Java) (Section 10.1). ",
     "https://opendatastructures.org/ods-java/10_1_BinaryHeap_Implicit_Bi.html"),
    ("Morin, P. (n.d.-b). ChainedHashTable: Hashing with chaining. In Open data structures (in Java) (Section 5.1). ",
     "https://opendatastructures.org/ods-java/5_1_ChainedHashTable_Hashin.html"),
    ("Python Software Foundation. (n.d.). Time: Time access and conversions. In Python 3.14 documentation. ",
     "https://docs.python.org/3.14/library/time.html#time.perf_counter"),
]


def benchmark_data():
    source = ROOT / "results/performance_results.txt"
    raw = source.read_text(encoding="utf-8")
    rows = re.findall(
        r"^\s*(\d+)\s+(Beginning|Early|Middle|Late|End|Missing)\s+(yes|no)\s+"
        r"([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)\s+([\d.]+)x\s*$",
        raw, re.MULTILINE,
    )
    if [int(row[0]) for row in rows] != [1001, 1025, 1075, 1125, 1150, 9999]:
        raise ValueError("saved benchmark does not contain the six required targets")
    for required in ("Dataset records: 150", "Hash table buckets: 211",
                     "Warm-up searches per target per algorithm: 100",
                     "Measured repetitions per target per algorithm: 10000"):
        if required not in raw:
            raise ValueError(f"unexpected benchmark configuration: {required}")
    def value(label):
        return re.search(rf"^{re.escape(label)}: ([\d.]+)", raw, re.MULTILINE)[1]
    return {
        "rows": rows,
        "hash_total": value("Overall hash total (s)"),
        "linear_total": value("Overall linear total (s)"),
        "hash_average": value("Overall average hash search (us)"),
        "linear_average": value("Overall average linear search (us)"),
        "ratio": value("Overall linear/hash time ratio"),
        "source": "results/performance_results.txt",
        "sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
    }


def bookmark(paragraph, name, identifier):
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(identifier))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(identifier))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def build(author=None, instructor=None, date="October 18, 2026"):
    data = benchmark_data()
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = Inches(0.5)
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = "Times New Roman", Pt(12)
    # Exact 24-point leading gives consistent double spacing for 12-point
    # text across Word and the PDF renderer.
    normal.paragraph_format.line_spacing = Pt(24)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.widow_control = True
    for style_name in ("Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name, style.font.size, style.font.bold = "Times New Roman", Pt(12), True
        style.font.color.rgb = normal.font.color.rgb
        style.paragraph_format.line_spacing = Pt(24)
        style.paragraph_format.space_before = style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.keep_with_next = True
    doc.styles["Heading 1"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Inches(0)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    header._p.append(field)
    doc.core_properties.title = TITLE
    doc.core_properties.author = author or ""
    doc.core_properties.subject = "CSC506: Design and Analysis of Algorithms"
    markdown = [f"# {TITLE}", "", *([author, ""] if author else []),
                "Colorado State University Global", "",
                "CSC506: Design and Analysis of Algorithms", "",
                *([instructor, ""] if instructor else []), date, ""]

    def paragraph(text, *, indent=True, bold=False, italic=False):
        p = doc.add_paragraph()
        if not indent:
            p.paragraph_format.first_line_indent = Inches(0)
        run = p.add_run(text)
        run.bold, run.italic = bold, italic
        markdown.extend([text, ""])
        return p

    def page(title, mark=None):
        doc.add_page_break()
        p = doc.add_paragraph(title, style="Heading 1")
        markdown.extend([f"## {title}", ""])
        if mark:
            bookmark(p, mark, len(markdown))

    def figure(number, title, filename, note):
        paragraph(f"Figure {number}", indent=False, bold=True)
        paragraph(title, indent=False, italic=True)
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1
        p.add_run().add_picture(str(ROOT / "docs/diagrams" / filename), width=Inches(6.5))
        paragraph(f"Note. {note}", indent=False)
        markdown.extend([f"![{title}](diagrams/{filename})", ""])

    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Inches(0)
    p = paragraph(TITLE, indent=False, bold=True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()
    for line in ([author] if author else []) + ["Colorado State University Global",
            "CSC506: Design and Analysis of Algorithms"] + ([instructor] if instructor else []) + [date]:
        p = doc.add_paragraph(line)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
    # The title-page helper above also appended its title once; normalize the
    # readable companion to the single title already recorded at its beginning.
    markdown = markdown[:-2]

    page("Project Design and Architecture")
    paragraph("HeapHash Nexus implements a custom hash table and a binary max-heap priority queue over the same 150 student records. Each immutable record contains a student ID, name, program, and integer priority. The supplied CSV uses IDs 1001–1150; validation permits priorities 1–100. The goal is correct keyed retrieval and priority scheduling, followed by a measured comparison with manual linear search.")
    figure(1, "Project Architecture", "project_architecture.png",
           "The CLI coordinates reusable modules. Algorithms remain outside menus.")

    page("Dataset and Data Flow")
    paragraph("The generator uses seed 42 for reproducible synthetic student data. CSV loading checks schema, field types, duplicate IDs, count, and priority bounds before exposing records. Application state retains an original snapshot and creates independent hash and heap containers; immutable records can safely be shared. CLI edits affect those containers, while reset reloads the CSV atomically. Demonstrations use isolated containers, and benchmarks verify agreement before measuring.")
    figure(2, "Dataset, Operations, and Results", "data_flow.png",
           "Benchmark output is separate from the source dataset. Reset restores original records.")

    page("Hash Function Selection", "HashExplanationStart")
    paragraph("The hash table maps an integer student ID to a bucket using h(k) = k mod m. Its normal configuration sets m = 211, so each result is an index from 0 through 210. For example, 1001 mod 211 = 157. The helper rejects non-integer keys, including Boolean values, and rejects invalid table sizes. Insertion, search, and deletion all call the same helper, ensuring that a key always selects the same chain.")
    paragraph("The table owns 211 distinct Python lists rather than using a dictionary as its key-value storage. Each list holds explicit (key, value) pairs. The prime size leaves room beyond the 150-record dataset and gives a load factor of α = n/m = 150/211 ≈ 0.7109. A prime modulus alone cannot guarantee an even distribution: IDs separated by 211 still collide. Because these particular 150 IDs are consecutive and span less than 211, their residues are distinct; the normal dataset therefore has no occupied chain longer than one.")
    paragraph("Chaining analysis depends on both bucket distribution and load, rather than the word “hash” alone. Under a suitable distribution, expected lookup work is O(1 + α); if every key lands in one bucket, lookup becomes O(n) (Morin, n.d.-b). This project uses a fixed modulus and does not resize or randomize its hash function. Its fast normal-dataset searches should therefore be interpreted with the actual key pattern in mind.")
    paragraph("This simple arithmetic function supports reproducibility, easy auditing, and deliberate collision examples. Statistics expose item count, load factor, occupied buckets, and chain lengths so that distribution can be inspected rather than assumed. The separate size-10 collision demonstration tests behavior that the normal size-211 dataset does not exercise.")

    page("Collision Resolution", "CollisionExplanationStart")
    paragraph("Separate chaining keeps records with equal hash indices together in one bucket while preserving their distinct keys. An operation computes the index and scans only that list. Insertion replaces a pair when its key already exists; otherwise it appends a new pair and increments item_count. Updating a duplicate does not increase the count. Search compares keys and returns the stored value, or None when no match exists.")
    paragraph("Deletion removes only the matching pair and decrements item_count once. It returns True for a removal and False for a missing key, preserving the count after unsuccessful deletion. Empty buckets are valid lists. Public bucket snapshots are immutable tuples, preventing callers from bypassing the methods and corrupting the count. The implementation does not require tombstones because deletion does not disrupt a probe sequence.")
    paragraph("The deliberate collision case chooses m = 10. IDs 1001, 1011, and 1021 all yield index 1, producing a three-record chain. Tests find every record, delete the middle key 1011, then find 1001 and 1021 again. Additional tests exercise duplicate updates, missing keys, collision-heavy operations, and count consistency. These checks establish that a shared index does not imply record replacement or accidental deletion of neighboring records.")
    paragraph("Chaining trades additional list storage and comparisons for straightforward insertion and deletion. An overloaded or poorly distributed fixed table can develop long chains; increasing capacity or rehashing would be future extensions. The current educational implementation instead exposes that limitation directly. Bucket storage is O(m + n), and an operation can require O(n) work in the worst case, including movement within a Python list after deletion. Correctness is independent of whether collisions are rare or frequent.")
    bookmark(doc.paragraphs[-1], "HashExplanationEnd", 999)

    page("Perfect and Non-Perfect Hashing")
    paragraph("A perfect hash function maps a specified key set to distinct indices without collisions (Black, 2019). Minimal perfect hashing additionally uses exactly as many output positions as keys. For the fixed keys 1001–1150, h(k) = k − 1001 is a minimal perfect mapping into 150 positions. This construction relies on the known, contiguous set; it is not a general solution for arbitrary student IDs.")
    paragraph("The implemented modulo-211 function is also collision-free on the supplied finite set, so it is perfect relative to that set, although it is not minimal. The distinction matters: a collision-free observation does not provide a guarantee for future keys. IDs 1001 and 1212 both map to index 157. Supporting arbitrary integer IDs requires the implemented chaining behavior even though the original CSV does not trigger it.")
    paragraph("A non-perfect mapping allows collisions and needs a resolution strategy. Separate chaining meets that requirement here without specializing the API to a frozen roster. A static perfect mapping could reduce comparisons for a known collection, but changing the set may require reconstruction or additional handling. The project prioritizes transparent operations, configurable capacity, and verified collision behavior over construction of a specialized perfect-hashing scheme. Neither the benchmark nor the code claims collision-free behavior for all possible inputs.")

    page("Binary Max-Heap Priority Queue")
    paragraph("The priority queue stores a complete binary tree in a Python list, with the maximum-priority record at index 0. For index i, the parent is (i − 1) // 2, and children are 2i + 1 and 2i + 2. Array-based binary heaps use these relationships to avoid explicit node pointers (Morin, n.d.-a). This project implements a max-heap by requiring each parent priority to be at least each child priority; it does not use heapq.")
    paragraph("Insertion rejects a duplicate student ID, appends the new record, and swaps it upward while its priority exceeds its parent’s. Peek returns the root without mutation. Extract-max replaces the removed root with the final record, then repeatedly swaps with the higher-priority child. Search scans IDs. Arbitrary deletion first locates the ID and moves the final record into the gap, then chooses upward repair if it exceeds its new parent; otherwise it repairs downward.")
    paragraph("Heap repair is O(log n), but public insertion is O(n) because duplicate-ID checking scans the heap. Search and arbitrary deletion are also O(n); extract-max is O(log n), peek is O(1), and validation is O(n). Storage is O(n). These bounds describe the actual implementation, including validation work, rather than assigning an idealized insertion bound to the complete public method.")
    paragraph("Equal priorities satisfy the invariant and need not retain insertion order. Empty peek and extraction return None; missing deletion returns False. Immutable snapshots prevent external list mutation. Tests validate the heap after insertion, deletion, and extraction, cover both repair directions, and confirm that repeated extraction yields nonincreasing priorities across all 150 records. Scheduling changes remain independent of the hash table.")

    page("Testing and User Interface")
    paragraph("The final automated suite contains 312 passing tests: 211 unit tests and 101 integration tests. Unit coverage includes dataset generation and validation (74), hash-table behavior (49), priority-queue behavior (71), and reusable CLI input validation (17). Integration coverage includes shared-dataset operations (15), performance correctness and result reporting (58), and CLI workflows (28). No timing threshold is required for a test to pass.")
    paragraph("Integration tests load all 150 records into both structures, compare searches with expected records, delete representative IDs, verify counts, and confirm independent state. Collision tests use a deliberately small table. Heap tests cover ties, empty cases, missing keys, and arbitrary deletion. Performance tests check all-record search agreement, reject mismatched tables, and use synthetic clocks only to verify timing arithmetic and execution protocol; the reported measurements come from a separate real run.")
    paragraph("The CLI exposes Hash Table, Priority Queue, Performance, and Dataset menus, a complete demonstration, reset, and exit. Reusable prompts reject invalid integers and out-of-range choices. Menus delegate to the core methods instead of implementing algorithms. Reset builds replacement state before assignment, so a failed reload preserves existing state. EOF and keyboard interruption exit gracefully. The noninteractive command python main.py --demo provides a reproducible end-to-end check.")
    paragraph("The recorded final demonstration passed dataset validation, ordinary hash operations, deliberate collisions and deletion, queue operations, peek, extract-max, heap validation, independent state, benchmark agreement, and result saving. Final regression output and a requirement-to-evidence matrix accompany the report. These checks establish functional behavior for the tested inputs; they do not constitute a proof that every possible input or environment has been exhaustively tested.")

    page("Benchmark Method")
    paragraph("The comparison uses the original ordered 150-record dataset and a normal 211-bucket hash table. The baseline manually scans records until it finds the student ID or reaches the end. Required targets are 1001, 1025, 1075, 1125, 1150, and missing ID 9999. These correspond to positions 1, 25, 75, 125, and 150, plus a full unsuccessful scan. Hash and linear search must return matching records before measurement begins.")
    paragraph("For each target and algorithm, the runner performs 100 warm-up searches, followed by 10,000 measured searches. Thus each algorithm executes 60,000 measured searches overall. Setup, agreement checks, warm-up, printing, and file writing are outside timed regions. The first measured algorithm alternates between targets to reduce a fixed ordering bias. Each measured block includes the loop and public search calls, so it reflects their combined overhead.")
    paragraph("Elapsed intervals use time.perf_counter(), a high-resolution elapsed-time counter intended for short-duration measurements (Python Software Foundation, n.d.). The saved final run identifies Python 3.14.4 and Windows-11-10.0.26300-SP0. Average time is total elapsed seconds divided by repetitions, converted to microseconds. The linear/hash ratio divides the linear total by the hash total; values above one indicate lower hash-search time.")
    paragraph("The aggregate ratio divides summed linear time by summed hash time, rather than averaging individual ratios. A zero hash total would be reported as unavailable. Results are stored in results/performance_results.txt, and the report builder records that file’s SHA-256 digest so the paper can be checked against its source. The single-machine experiment measures one block per method and target; it does not estimate confidence intervals or performance at larger dataset sizes.")

    page("Observed Performance")
    paragraph("Table 1 presents the actual saved final benchmark. All five present IDs were found by both methods, and both returned None for 9999. Time values are arithmetic averages per search, not single-call timer readings.")
    paragraph("Table 1", indent=False, bold=True)
    paragraph("Final Hash-Table and Linear-Search Measurements", indent=False, italic=True)
    table = doc.add_table(rows=1, cols=5)
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "bottom"):
        rule = OxmlElement(f"w:{edge}")
        rule.set(qn("w:val"), "single")
        rule.set(qn("w:sz"), "6")
        borders.append(rule)
    table._tbl.tblPr.append(borders)
    headers = ["Student ID", "Scenario", "Hash (μs)", "Linear (μs)", "Linear/hash"]
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    markdown.extend(["| " + " | ".join(headers) + " |", "|---|---|---:|---:|---:|"])
    for row in data["rows"]:
        values = [row[0], row[1], row[5], row[6], row[7] + "×"]
        for cell, value in zip(table.add_row().cells, values):
            cell.text = value
        markdown.append("| " + " | ".join(values) + " |")
    markdown.append("")
    for index, row in enumerate(table.rows):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.first_line_indent = Inches(0)
                p.paragraph_format.line_spacing = 1.25
                p.paragraph_format.space_after = Pt(5)
                for run in p.runs:
                    run.font.name, run.font.size, run.bold = "Times New Roman", Pt(11), index == 0
    paragraph("Note. Each row uses 10,000 measured searches per algorithm after 100 warm-ups. Ratios above one favor hash search; a ratio below one favors linear search. Values retain the saved report’s rounding.", indent=False)
    paragraph(f"Across all 60,000 searches per method, hash search took {data['hash_total']} s and linear search took {data['linear_total']} s. Overall averages were {data['hash_average']} μs and {data['linear_average']} μs, respectively, yielding a {data['ratio']}× aggregate linear/hash ratio. The first-record linear search was faster, whereas hash search was faster for the remaining targets in this run.")
    paragraph("Later and missing linear searches examine more records, while normal-table hash searches inspect at most one stored key. The results are consistent with that difference, but the data are small, ordered, synthetic, and collision-free under modulus 211. Interpreter overhead, scheduling, and machine state affect submicrosecond averages. These measurements support this dataset comparison and do not establish a universal speedup or characterize collision-heavy performance.")

    page("Conclusion")
    paragraph("The completed project demonstrates two independently useful structures over one validated dataset. The hash table provides keyed updates and retrieval with explicit collision chains. The max-heap maintains a highest-priority root while supporting insertion, lookup, deletion, peek, and extraction. Clear module boundaries keep the CLI focused on input and presentation, and immutable records support safe sharing between independent containers.")
    paragraph(f"All 312 automated tests and the final complete demonstration passed. The real benchmark recorded a {data['ratio']}× aggregate linear/hash time ratio across the specified six targets, while showing that a first-element linear lookup can outperform hashing. The conclusion therefore depends on workload and data distribution rather than claiming an unconditional advantage. The collision demonstration and heap-validation tests complement the normal-dataset benchmark by checking structural cases that timings alone cannot establish.")
    paragraph("Future work could add hash-table resizing, repeated independent benchmark blocks with uncertainty estimates, larger and less uniform datasets, and an ID index for faster priority-queue lookup. Such extensions would change the current design and should be measured and tested separately. The submitted source, documented commands, diagrams, saved measurements, regression evidence, and traceability matrix allow the implemented behavior to be reviewed and reproduced.")

    page("References")
    for citation, url in REFERENCES:
        p = paragraph(citation + url, indent=False)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        # APA uses italic book/document titles in the reference list.
        text = p.text
        if citation.startswith("Black"):
            italic_text = "Dictionary of Algorithms and Data Structures"
        elif citation.startswith("Morin"):
            italic_text = "Open data structures (in Java)"
        else:
            italic_text = "Python 3.14 documentation"
        before, after = text.split(italic_text, 1)
        p.clear()
        p.add_run(before)
        p.add_run(italic_text).italic = True
        p.add_run(after)

    destination = ROOT / "docs"
    doc.save(destination / "final_report.docx")
    (destination / "final_report.md").write_text("\n".join(markdown), encoding="utf-8")
    metadata = {"title": TITLE, "date": date, "author": author, "instructor": instructor,
                "benchmark": data, "test_count": 312,
                "hash_explanation_bookmarks": ["HashExplanationStart", "CollisionExplanationStart", "HashExplanationEnd"]}
    (destination / "report_metadata.json").write_text(json.dumps(metadata, indent=2) + "\n", encoding="utf-8")
    print("Created docs/final_report.docx, docs/final_report.md, and docs/report_metadata.json")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--author")
    parser.add_argument("--instructor")
    parser.add_argument("--date", default="October 18, 2026")
    args = parser.parse_args()
    build(args.author, args.instructor, args.date)
